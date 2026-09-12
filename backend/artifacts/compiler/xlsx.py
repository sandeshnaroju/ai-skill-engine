"""
backend/artifacts/compiler/xlsx.py
Compiles spreadsheet artifacts into beautifully formatted Microsoft Excel (.xlsx)
with styled headers, zebra rows, auto-sized columns, and freeze panes.
"""
import io
import re
import json
from typing import List, Any, Dict

from models import SessionArtifact
from .utils import _strip_html


def _parse_spreadsheet_blocks(blocks: List[Any], default_title: str = "Sheet1") -> List[Dict[str, Any]]:
    """
    Parses artifact blocks in any format (JSON, Markdown Table, CSV, TSV) into a list of sheet definitions:
    [{"name": "Sheet 1", "columns": ["Col A", "Col B"], "rows": [["val1", "val2"]]}]
    """
    import csv as pycsv
    sheets = []

    for idx, block in enumerate(blocks):
        raw = str(getattr(block, 'content', '') or '').strip()
        b_title = getattr(block, 'title', '') or f"Sheet {idx + 1}"
        if not raw:
            continue

        # Strip markdown code fences if present (```csv ... ```, ```json ... ```, etc.)
        cleaned = re.sub(r'^```(?:[a-zA-Z0-9_-]+)?\s*\n', '', raw)
        cleaned = re.sub(r'\n```\s*$', '', cleaned).strip()

        # 1. Try parsing JSON format
        try:
            data = json.loads(cleaned)
            if isinstance(data, dict):
                # { "sheets": [ { "sheet_name": "...", "columns": [...], "rows": [...] }, ... ] }
                if "sheets" in data and isinstance(data["sheets"], list) and len(data["sheets"]) > 0:
                    for s_idx, s in enumerate(data["sheets"]):
                        s_name = s.get("sheet_name") or s.get("name") or f"{b_title} {s_idx + 1}"
                        cols = s.get("columns") or []
                        rows = s.get("rows") or []
                        if cols or rows:
                            sheets.append({"name": s_name[:31], "columns": cols, "rows": rows})
                    continue

                # { "sheet_name": "...", "columns": [...], "rows": [...] }
                s_name = data.get("sheet_name") or data.get("name") or b_title
                cols = data.get("columns") or []
                rows = data.get("rows") or []
                if cols or rows:
                    sheets.append({"name": s_name[:31], "columns": cols, "rows": rows})
                    continue

            elif isinstance(data, list) and len(data) > 0:
                # [ { "Col A": "Val 1", "Col B": "Val 2" }, ... ]
                if isinstance(data[0], dict):
                    headers = list(data[0].keys())
                    rows = [[row.get(h, "") for h in headers] for row in data]
                    sheets.append({"name": b_title[:31], "columns": headers, "rows": rows})
                    continue
                # [ ["Col A", "Col B"], ["Val 1", "Val 2"] ]
                elif isinstance(data[0], list):
                    headers = data[0]
                    rows = data[1:]
                    sheets.append({"name": b_title[:31], "columns": headers, "rows": rows})
                    continue
        except Exception:
            pass

        # 2. Try parsing Markdown Table format (| Col 1 | Col 2 |)
        if "|" in cleaned:
            table_lines = [l.strip() for l in cleaned.splitlines() if l.strip().startswith("|") and l.strip().endswith("|")]
            if len(table_lines) >= 2:
                def parse_md_row(line):
                    return [c.strip() for c in line.strip("|").split("|")]

                headers = parse_md_row(table_lines[0])
                data_lines = table_lines[1:]
                # Exclude divider row (e.g. |---|---|)
                if data_lines and all(re.match(r'^:?-+:?$', c) for c in parse_md_row(data_lines[0]) if c):
                    data_lines = data_lines[1:]

                rows = [parse_md_row(l) for l in data_lines]
                sheets.append({"name": b_title[:31], "columns": headers, "rows": rows})
                continue

        # 3. Try CSV / TSV format
        try:
            delimiter = '\t' if '\t' in cleaned else ','
            reader = pycsv.reader(cleaned.splitlines(), delimiter=delimiter)
            all_rows = [r for r in reader if any(field.strip() for field in r)]
            if all_rows:
                headers = [h.strip() for h in all_rows[0]]
                rows = [[c.strip() for c in r] for r in all_rows[1:]]
                sheets.append({"name": b_title[:31], "columns": headers, "rows": rows})
                continue
        except Exception:
            pass

        # 4. Fallback: line-by-line raw text
        lines = [l.strip() for l in cleaned.splitlines() if l.strip()]
        sheets.append({"name": b_title[:31], "columns": ["Content"], "rows": [[l] for l in lines]})

    if not sheets:
        sheets.append({"name": default_title[:31], "columns": ["Column 1", "Column 2"], "rows": [["", ""]]})

    return sheets


def compile_to_xlsx(artifact: SessionArtifact) -> io.BytesIO:
    """Compile spreadsheet into beautifully formatted Microsoft Excel (.xlsx) with styled headers, zebra rows, auto-sized columns, and freeze panes."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # Remove initial blank sheet

    blocks = sorted(artifact.blocks, key=lambda b: b.order_index) if artifact.blocks else []
    parsed_sheets = _parse_spreadsheet_blocks(blocks, default_title=artifact.title or "Sheet1")

    # Styling definitions
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")  # Dark slate
    data_font = Font(name="Calibri", size=10, color="0F172A")
    alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")  # Subtle zebra
    white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    for s_data in parsed_sheets:
        ws = wb.create_sheet(title=s_data["name"][:31])
        ws.views.sheetView[0].showGridLines = True

        columns = s_data.get("columns") or []
        rows = s_data.get("rows") or []

        # If columns is empty but rows exist, determine max columns
        if not columns and rows:
            max_c = max(len(r) for r in rows)
            columns = [f"Col {get_column_letter(i+1)}" for i in range(max_c)]

        num_cols = len(columns)
        col_max_lengths = [max(len(str(c)), 8) for c in columns]

        # 1. Render Header Row
        ws.row_dimensions[1].height = 26.0
        for col_idx, col_name in enumerate(columns, 1):
            cell = ws.cell(row=1, column=col_idx, value=_strip_html(str(col_name)))
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin_border

        # 2. Render Data Rows
        for row_idx, row_values in enumerate(rows, 2):
            ws.row_dimensions[row_idx].height = 20.0
            row_fill = alt_fill if (row_idx % 2 == 1) else white_fill

            for col_idx in range(1, num_cols + 1):
                val = row_values[col_idx - 1] if (col_idx - 1) < len(row_values) else ""
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.font = data_font
                cell.fill = row_fill
                cell.border = thin_border

                # Value formatting & auto-typing
                if val is None or val == "":
                    cell.value = ""
                    cell.alignment = Alignment(vertical="center")
                elif isinstance(val, (int, float)):
                    cell.value = val
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                    if isinstance(val, float):
                        cell.number_format = '#,##0.00'
                else:
                    str_val = _strip_html(str(val)).strip()
                    # Check for integer
                    if re.match(r'^-?\d+$', str_val):
                        try:
                            cell.value = int(str_val)
                            cell.alignment = Alignment(horizontal="right", vertical="center")
                            cell.number_format = '#,##0'
                        except ValueError:
                            cell.value = str_val
                            cell.alignment = Alignment(horizontal="left", vertical="center")
                    # Check for float
                    elif re.match(r'^-?\d+\.\d+$', str_val):
                        try:
                            cell.value = float(str_val)
                            cell.alignment = Alignment(horizontal="right", vertical="center")
                            cell.number_format = '#,##0.00'
                        except ValueError:
                            cell.value = str_val
                            cell.alignment = Alignment(horizontal="left", vertical="center")
                    # Check for Currency (e.g. $1,234.56 or -$50.00)
                    elif re.match(r'^-?\$?\s*([0-9]{1,3}(?:,[0-9]{3})*(?:\.[0-9]+)?|\.[0-9]+)\$?$', str_val):
                        try:
                            clean_num = re.sub(r'[\$,\s]', '', str_val)
                            cell.value = float(clean_num)
                            cell.alignment = Alignment(horizontal="right", vertical="center")
                            cell.number_format = '$#,##0.00'
                        except ValueError:
                            cell.value = str_val
                            cell.alignment = Alignment(horizontal="left", vertical="center")
                    # Check for Percentage (e.g. 25.5%)
                    elif re.match(r'^-?\d+(?:\.\d+)?%$', str_val):
                        try:
                            clean_pct = str_val.rstrip('%')
                            cell.value = float(clean_pct) / 100.0
                            cell.alignment = Alignment(horizontal="right", vertical="center")
                            cell.number_format = '0.0%'
                        except ValueError:
                            cell.value = str_val
                            cell.alignment = Alignment(horizontal="left", vertical="center")
                    else:
                        cell.value = str_val
                        cell.alignment = Alignment(horizontal="left", vertical="center")

                # Track column width
                str_len = len(str(cell.value or ''))
                if (col_idx - 1) < len(col_max_lengths):
                    col_max_lengths[col_idx - 1] = max(col_max_lengths[col_idx - 1], str_len)

        # 3. Freeze top header row so it stays visible during scrolling
        ws.freeze_panes = "A2"

        # 4. Auto-fit column widths
        for col_idx, max_l in enumerate(col_max_lengths, 1):
            col_letter = get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = min(max(max_l + 4, 12), 48)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
