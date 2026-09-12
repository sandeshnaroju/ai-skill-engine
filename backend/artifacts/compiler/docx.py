"""
backend/artifacts/compiler/docx.py
Compiles Markdown / Document artifacts into native Microsoft Word (.docx)
with styled tables, headings, code blocks, hyperlinks, images, and horizontal dividers.
"""
import io
import re
from typing import Any

from models import SessionArtifact
from .utils import _strip_html, _tokenize_inline_formatting, _fetch_image_bytes, _hex_to_rgb


def compile_to_docx(artifact: SessionArtifact) -> io.BytesIO:
    """Compile Markdown/Document artifact into native Microsoft Word (.docx) with real tables, hyperlinks, code blocks, and styled typography."""
    import docx
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    doc = docx.Document()

    # Set 1-inch margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Document Title
    title_p = doc.add_paragraph()
    title_run = title_p.add_run(artifact.title)
    title_run.font.name = "Calibri"
    title_run.font.size = Pt(24)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(15, 23, 42)
    title_p.paragraph_format.space_after = Pt(16)

    def add_hyperlink(paragraph, url, text, color="2563EB"):
        """Add a real clickable external hyperlink to a python-docx paragraph."""
        try:
            part = paragraph.part
            r_id = part.relate_to(url, docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
            hyperlink = parse_xml(f'<w:hyperlink {nsdecls("w")} r:id="{r_id}" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"/>')
            new_run = parse_xml(f'<w:r {nsdecls("w")}><w:rPr><w:color w:val="{color}"/><w:u w:val="single"/><w:rFonts w:ascii="Calibri" w:hAnsi="Calibri"/></w:rPr><w:t>{_strip_html(text)}</w:t></w:r>')
            hyperlink.append(new_run)
            paragraph._p.append(hyperlink)
        except Exception:
            run = paragraph.add_run(_strip_html(text))
            run.font.name = "Calibri"
            run.font.color.rgb = RGBColor(37, 99, 235)
            run.font.underline = True

    def apply_inline_tokens_to_paragraph(p, text_content):
        tokens = _tokenize_inline_formatting(text_content)
        if not tokens:
            return
        for t in tokens:
            t_type = t.get("type", "text")
            t_text = _strip_html(t.get("text", ""))
            if not t_text:
                continue

            if t_type == "link":
                add_hyperlink(p, t.get("url", "#"), t_text)
                continue

            run = p.add_run(t_text)
            run.font.name = "Calibri"
            run.font.size = Pt(11)

            if t_type == "bold":
                run.bold = True
                run.font.color.rgb = RGBColor(15, 23, 42)
            elif t_type == "italic":
                run.italic = True
                run.font.color.rgb = RGBColor(51, 65, 85)
            elif t_type == "code":
                run.font.name = "Consolas"
                run.font.size = Pt(10)
                run.font.color.rgb = RGBColor(124, 58, 237)
            elif t_type == "colored":
                r, g, b = _hex_to_rgb(t.get("color", "#7c3aed"))
                run.font.color.rgb = RGBColor(r, g, b)
                run.bold = True
            else:
                run.font.color.rgb = RGBColor(51, 65, 85)

    def set_cell_shading(cell, color_hex):
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex.lstrip("#")}"/>')
        cell._tc.get_or_add_tcPr().append(shading)

    def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
        tcPr.append(tcMar)

    def add_horizontal_rule():
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(8)
        pPr = p._p.get_or_add_pPr()
        pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="1" w:color="CBD5E1"/></w:pBdr>')
        pPr.append(pBdr)

    def add_code_block_table(code_text):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tblPr = tbl._tbl.tblPr
        borders = parse_xml(f'<w:tblBorders {nsdecls("w")}><w:top w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/><w:bottom w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/><w:left w:val="single" w:sz="12" w:space="0" w:color="6366F1"/><w:right w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/></w:tblBorders>')
        tblPr.append(borders)

        cell = tbl.cell(0, 0)
        set_cell_shading(cell, "F8FAFC")
        set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(code_text)
        run.font.name = "Consolas"
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(30, 41, 59)

        spacer = doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(6)

    def add_image_to_document(url_or_data, alt_text=""):
        img_bytes = _fetch_image_bytes(url_or_data)
        if img_bytes:
            try:
                from PIL import Image as PILImage
                pil_img = PILImage.open(io.BytesIO(img_bytes))
                w, h = pil_img.size
                max_width_in = 5.5
                aspect = h / w if w > 0 else 0.75
                img_w_in = min(w / 96.0, max_width_in)
                if img_w_in < 1.0:
                    img_w_in = max_width_in
                img_h_in = img_w_in * aspect
                if img_h_in > 6.0:
                    img_h_in = 6.0
                    img_w_in = img_h_in / aspect

                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(6)
                p.paragraph_format.space_after = Pt(3)
                run = p.add_run()
                run.add_picture(io.BytesIO(img_bytes), width=Inches(img_w_in))

                if alt_text:
                    caption_p = doc.add_paragraph()
                    caption_p.paragraph_format.space_after = Pt(8)
                    caption_run = caption_p.add_run(f"Figure: {_strip_html(alt_text)}")
                    caption_run.font.name = "Calibri"
                    caption_run.font.size = Pt(9)
                    caption_run.font.italic = True
                    caption_run.font.color.rgb = RGBColor(100, 116, 139)
                return
            except Exception:
                pass

        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(f"🖼 [Image: {_strip_html(alt_text or 'Image')}]")
        run.font.name = "Calibri"
        run.font.size = Pt(10)
        run.font.italic = True
        run.font.color.rgb = RGBColor(100, 116, 139)

    blocks = sorted(artifact.blocks, key=lambda b: b.order_index) if artifact.blocks else []
    for block in blocks:
        lines = block.content.splitlines()
        in_code_block = False
        code_block_lines = []
        table_rows = []

        def flush_table():
            nonlocal table_rows
            if not table_rows:
                return
            clean_rows = []
            for row in table_rows:
                if all(re.match(r'^:?-+:?$', c.strip()) for c in row if c.strip()):
                    continue
                clean_rows.append(row)

            if clean_rows:
                cols_count = max(len(r) for r in clean_rows)
                tbl = doc.add_table(rows=len(clean_rows), cols=cols_count)
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

                tblPr = tbl._tbl.tblPr
                borders = parse_xml(f'<w:tblBorders {nsdecls("w")}><w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/><w:bottom w:val="single" w:sz="6" w:space="0" w:color="94A3B8"/><w:left w:val="none"/><w:right w:val="none"/><w:insideH w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/><w:insideV w:val="none"/></w:tblBorders>')
                tblPr.append(borders)

                for r_idx, row_data in enumerate(clean_rows):
                    is_header = (r_idx == 0)
                    for c_idx in range(cols_count):
                        cell_val = row_data[c_idx] if c_idx < len(row_data) else ""
                        cell = tbl.cell(r_idx, c_idx)
                        set_cell_margins(cell, top=120, bottom=120, left=180, right=180)

                        if is_header:
                            set_cell_shading(cell, "1E293B")
                            p = cell.paragraphs[0]
                            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                            run = p.add_run(_strip_html(cell_val))
                            run.font.name = "Calibri"
                            run.font.size = Pt(10)
                            run.font.bold = True
                            run.font.color.rgb = RGBColor(255, 255, 255)
                        else:
                            if r_idx % 2 == 1:
                                set_cell_shading(cell, "F8FAFC")
                            else:
                                set_cell_shading(cell, "FFFFFF")
                            p = cell.paragraphs[0]
                            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                            apply_inline_tokens_to_paragraph(p, cell_val)

                spacer_p = doc.add_paragraph()
                spacer_p.paragraph_format.space_after = Pt(8)

            table_rows = []

        for line in lines:
            raw_line = line
            line_str = line.strip()

            # Handle Code Blocks
            if line_str.startswith("```"):
                if in_code_block:
                    in_code_block = False
                    code_text = "\n".join(code_block_lines)
                    add_code_block_table(code_text)
                    code_block_lines = []
                else:
                    flush_table()
                    in_code_block = True
                    code_block_lines = []
                continue

            if in_code_block:
                code_block_lines.append(raw_line)
                continue

            # Handle Tables
            if line_str.startswith("|") and line_str.endswith("|"):
                cells = [c.strip() for c in line_str.strip("|").split("|")]
                table_rows.append(cells)
                continue
            else:
                flush_table()

            if not line_str:
                continue

            # Horizontal Rule (---, ***, ___, <hr>)
            if re.match(r'^(?:---|___|\*\*\*|\-{3,}|\*{3,}|_{3,}|<hr\s*/?>)$', line_str):
                add_horizontal_rule()
                continue

            # Markdown Image: ![alt](url)
            img_match = re.match(r'^!\[(.*?)\]\((.*?)\)$', line_str)
            if img_match:
                alt_txt, img_url = img_match.group(1), img_match.group(2)
                add_image_to_document(img_url, alt_txt)
                continue

            # HTML Image: <img ... src="..." ...>
            html_img_match = re.search(r'<img\s+[^>]*src=[\'"]([^\'"]+)[\'"][^>]*>', line_str)
            if html_img_match:
                img_url = html_img_match.group(1)
                alt_match = re.search(r'alt=[\'"]([^\'"]+)[\'"]', line_str)
                alt_txt = alt_match.group(1) if alt_match else ""
                add_image_to_document(img_url, alt_txt)
                continue

            # Headings
            if line_str.startswith("# "):
                h = doc.add_paragraph()
                h.paragraph_format.space_before = Pt(16)
                h.paragraph_format.space_after = Pt(6)
                run = h.add_run(_strip_html(line_str[2:].strip()))
                run.font.name = "Calibri"
                run.font.size = Pt(18)
                run.font.bold = True
                run.font.color.rgb = RGBColor(30, 41, 59)
            elif line_str.startswith("## "):
                h = doc.add_paragraph()
                h.paragraph_format.space_before = Pt(12)
                h.paragraph_format.space_after = Pt(4)
                run = h.add_run(_strip_html(line_str[3:].strip()))
                run.font.name = "Calibri"
                run.font.size = Pt(14)
                run.font.bold = True
                run.font.color.rgb = RGBColor(51, 65, 85)
            elif line_str.startswith("### "):
                h = doc.add_paragraph()
                h.paragraph_format.space_before = Pt(10)
                h.paragraph_format.space_after = Pt(3)
                run = h.add_run(_strip_html(line_str[4:].strip()))
                run.font.name = "Calibri"
                run.font.size = Pt(12)
                run.font.bold = True
                run.font.color.rgb = RGBColor(71, 85, 105)
            elif line_str.startswith("- ") or line_str.startswith("* "):
                p = doc.add_paragraph(style='List Bullet')
                p.paragraph_format.space_after = Pt(3)
                apply_inline_tokens_to_paragraph(p, line_str[2:].strip())
            elif re.match(r'^\d+\.\s+', line_str):
                text_val = re.sub(r'^\d+\.\s+', '', line_str)
                p = doc.add_paragraph(style='List Number')
                p.paragraph_format.space_after = Pt(3)
                apply_inline_tokens_to_paragraph(p, text_val)
            elif line_str.startswith(">"):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.25)
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(6)
                clean_quote = re.sub(r'^>\s*(\[!NOTE\]|\[!TIP\]|\[!IMPORTANT\])?\s*', '', line_str)
                apply_inline_tokens_to_paragraph(p, clean_quote)
            else:
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(6)
                apply_inline_tokens_to_paragraph(p, line_str)

        flush_table()

    output = io.BytesIO()
    doc.save(output)
    output.seek(0)
    return output
